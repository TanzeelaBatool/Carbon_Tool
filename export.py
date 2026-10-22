"""Excel export: every result table plus native Excel charts, written from the calculated results."""
from __future__ import annotations

import io

import pandas as pd
from openpyxl.chart import BarChart, Reference
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

import engine as E

HEAD = PatternFill("solid", fgColor="1F3A5F")


def _write(ws_name, df, writer, startrow=0, title=None):
    if title:
        pd.DataFrame([[title]]).to_excel(writer, sheet_name=ws_name, startrow=startrow, index=False, header=False)
        writer.sheets[ws_name].cell(row=startrow + 1, column=1).font = Font(bold=True, size=12)
        startrow += 1
    df.to_excel(writer, sheet_name=ws_name, startrow=startrow, index=False)
    return startrow


def _style(ws):
    for cell in ws[1]:
        pass
    for col in ws.columns:
        width = max((len(str(c.value)) for c in col if c.value is not None), default=8)
        ws.column_dimensions[get_column_letter(col[0].column)].width = min(max(10, width + 2), 60)


def _header_style(ws, row):
    for c in ws[row]:
        if c.value is not None:
            c.fill = HEAD
            c.font = Font(bold=True, color="FFFFFF")
            c.alignment = Alignment(wrap_text=True, vertical="center")


def _bar(ws, title, cats_col, val_cols, header_row, n, anchor, ytitle="tCO2e"):
    ch = BarChart()
    ch.type = "col"
    ch.title = title
    ch.y_axis.title = ytitle
    for vc in val_cols:
        ch.add_data(Reference(ws, min_col=vc, min_row=header_row, max_row=header_row + n), titles_from_data=True)
    ch.set_categories(Reference(ws, min_col=cats_col, min_row=header_row + 1, max_row=header_row + n))
    ch.height, ch.width = 8.5, 16
    ws.add_chart(ch, anchor)


def build_excel(res: dict, data: pd.DataFrame, log: list, settings: E.Settings) -> bytes:
    buf = io.BytesIO()
    with pd.ExcelWriter(buf, engine="openpyxl") as w:
        # Findings
        pd.DataFrame({"Key findings (generated from the calculated numbers)": res["findings"]}).to_excel(w, sheet_name="Findings", index=False)
        # Settings used
        pd.DataFrame({"Setting": ["Metrics", "Sizes excluded from analysis", "Missing value treatment", "Duplicate rows removed", "IQR multiplier",
                                  "Enterprises analysed"],
                      "Value": [", ".join(settings.metrics), ", ".join(settings.exclude_sizes) or "none",
                                "counted as 0 (report convention)" if settings.missing_as_zero else "left out of each metric's statistics",
                                "yes" if settings.drop_duplicates else "no", settings.iqr_k, len(res["population"])]}).to_excel(w, sheet_name="Settings used", index=False)
        res["describe"].to_excel(w, sheet_name="Summary statistics", index=False)
        # scope table + chart
        sc = res["scopes"].drop(columns=["Dominant in (text)"])
        sc.to_excel(w, sheet_name="Scope-wise", index=False)
        row = len(sc) + 3
        for g, t in res["by_size_scopes"].items():
            _write("Scope-wise", t.drop(columns=["Dominant in (text)"]), w, startrow=row, title=f"Scope-wise analysis - {g} enterprises")
            row += len(t) + 4
        # by size
        res["by_size"].to_excel(w, sheet_name="By size", index=False)
        if len(res["ratios"]):
            _write("By size", res["ratios"], w, startrow=len(res["by_size"]) + 3, title="Ratio of average footprints")
        res["ranking"].drop(columns=["Reported total"], errors="ignore").to_excel(w, sheet_name="Ranking", index=False)
        flag = res["outliers"].copy()
        flag.to_excel(w, sheet_name="Outliers", index=False)
        _write("Outliers", pd.DataFrame(res["fences"], index=[0]).T.reset_index().rename(columns={"index": "Quantity", 0: "Value"}), w,
               startrow=len(flag) + 3, title=f"IQR fences (k = {settings.iqr_k})")
        for k, nm in (("by_sector", "By sector"), ("by_location", "By location"), ("by_process", "By process")):
            if k in res:
                res[k].to_excel(w, sheet_name=nm, index=False)
        res["reconcile"].to_excel(w, sheet_name="Reconciliation", index=False)
        pd.DataFrame(log or [{"Severity": "", "Check": "No issues found", "Enterprise": "", "Field": "", "Detail": ""}]).to_excel(w, sheet_name="Data quality log", index=False)
        data.to_excel(w, sheet_name="Clean data", index=False)

        wb = w.book
        for ws in wb.worksheets:
            _style(ws)
            _header_style(ws, 1)
        # native charts
        ws = wb["Scope-wise"]; _bar(ws, "Total emissions by scope (tCO2e)", 1, [2], 1, len(sc), "K2")
        ws = wb["By size"]; n = len(res["by_size"]); _bar(ws, "Average footprint by size (tCO2e)", 1, [3], 1, n, "J2")
        ws = wb["Ranking"]
        top = min(15, len(res["ranking"]))
        colT = list(res["ranking"].drop(columns=["Reported total"], errors="ignore").columns).index("Total") + 1
        ws_ch = BarChart(); ws_ch.type = "bar"; ws_ch.title = f"Top {top} enterprises by total footprint (tCO2e)"
        ws_ch.add_data(Reference(ws, min_col=colT, min_row=1, max_row=top + 1), titles_from_data=True)
        ws_ch.set_categories(Reference(ws, min_col=3, min_row=2, max_row=top + 1))
        ws_ch.y_axis.scaling.orientation = "minMax"; ws_ch.x_axis.scaling.orientation = "maxMin"
        ws_ch.height, ws_ch.width = 10, 18
        ws.add_chart(ws_ch, "N2")
        if "by_sector" in res and len(res["by_sector"]) > 1:
            ws = wb["By sector"]; n = len(res["by_sector"])
            cols = [i + 1 for i, c in enumerate(res["by_sector"].columns) if c.endswith("(tCO2e)") and c.split(" (")[0] in settings.metrics]
            _bar(ws, "Scope totals by sector (tCO2e)", 1, cols, 1, n, "L2")
    return buf.getvalue()
